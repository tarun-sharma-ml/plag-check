#!/usr/bin/env python3
"""
Plagiarism Checker
-------------------
Checks a document against web sources using semantic similarity
(sentence embeddings) rather than plain string/keyword matching.

How it works:
1. Splits the input document into overlapping chunks (a few sentences each).
2. For each chunk, searches the web (DuckDuckGo, no API key needed) for
   candidate matching pages.
3. Downloads and cleans the text of those candidate pages.
4. Computes semantic similarity (cosine similarity of sentence embeddings)
   between your chunk and passages from each candidate page.
5. Flags any chunk above a similarity threshold as "possible plagiarism"
   and reports the best-matching source + score.

No paid API key is required:
- Web search: ddgs / DuckDuckGo (free)
- Embeddings: sentence-transformers, running a small local model
  (downloads once, then works offline)

Install dependencies:
    pip install -r requirements.txt

Usage:
    python plagiarism_checker.py my_document.txt
    python plagiarism_checker.py my_document.txt --threshold 0.75 --results-per-chunk 5
    python plagiarism_checker.py my_document.txt --json report.json
"""

import argparse
import json
import os
import re
import sys
import time
from dataclasses import dataclass, field
from typing import List, Optional

import requests
from bs4 import BeautifulSoup


# ---- Lazy imports for heavy libs (so --help works fast) --------------------
def _lazy_imports():
    global SentenceTransformer, util, DDGS
    from sentence_transformers import SentenceTransformer, util
    try:
        from ddgs import DDGS
    except ImportError:
        from duckduckgo_search import DDGS


# ---------------------------------------------------------------------------
# Document loading (.txt, .docx, .pdf)
# ---------------------------------------------------------------------------

def load_document(path: str) -> str:
    """Read text from a .txt, .docx or .pdf file."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            sys.exit("Reading .docx needs python-docx: pip install python-docx")
        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            sys.exit("Reading .pdf needs pypdf: pip install pypdf")
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".doc":
        sys.exit("Old .doc files are not supported. Save it as .docx or .pdf first.")

    # Plain text: try UTF-8, then fall back to a common Windows encoding
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------

@dataclass
class ChunkResult:
    chunk_index: int
    chunk_text: str
    best_score: float = 0.0
    best_source_url: Optional[str] = None
    best_source_snippet: Optional[str] = None
    is_flagged: bool = False
    sources_checked: int = 0


@dataclass
class Report:
    total_chunks: int = 0            # chunks actually checked
    skipped_chunks: int = 0          # too short to check
    no_result_chunks: int = 0        # search returned nothing usable
    flagged_chunks: List[ChunkResult] = field(default_factory=list)
    all_chunks: List[ChunkResult] = field(default_factory=list)

    @property
    def overall_score(self) -> float:
        """Rough overall plagiarism percentage = flagged chunks / total chunks."""
        if self.total_chunks == 0:
            return 0.0
        return 100.0 * len(self.flagged_chunks) / self.total_chunks

    @property
    def coverage(self) -> float:
        """Percentage of checked chunks that had at least one web source compared."""
        if self.total_chunks == 0:
            return 0.0
        return 100.0 * (self.total_chunks - self.no_result_chunks) / self.total_chunks


# ---------------------------------------------------------------------------
# Text processing
# ---------------------------------------------------------------------------

def split_into_sentences(text: str) -> List[str]:
    """Simple sentence splitter (no nltk download needed)."""
    text = re.sub(r"\s+", " ", text).strip()
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", text)
    return [s.strip() for s in sentences if s.strip()]


def chunk_document(text: str, sentences_per_chunk: int = 3, overlap: int = 1) -> List[str]:
    """Group sentences into overlapping chunks for more robust matching."""
    sentences = split_into_sentences(text)
    if not sentences:
        return []

    chunks = []
    step = max(1, sentences_per_chunk - overlap)
    for i in range(0, len(sentences), step):
        chunk = " ".join(sentences[i:i + sentences_per_chunk])
        if chunk:
            chunks.append(chunk)
        if i + sentences_per_chunk >= len(sentences):
            break
    return chunks


def clean_web_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
        tag.decompose()
    text = soup.get_text(separator=" ")
    return re.sub(r"\s+", " ", text).strip()


_REF_HEADING = re.compile(
    r"^\s*(references|bibliography|works cited|sources)\s*:?\s*$", re.IGNORECASE | re.MULTILINE
)


def strip_references(text: str) -> str:
    """Cut the reference list off the end so citations are not checked as prose."""
    matches = list(_REF_HEADING.finditer(text))
    if matches and matches[-1].start() > len(text) * 0.5:
        return text[: matches[-1].start()]
    return text


def build_query(chunk: str, max_words: int = 14) -> str:
    """Use the longest (most distinctive) sentence of the chunk, trimmed."""
    sentences = split_into_sentences(chunk) or [chunk]
    best = max(sentences, key=lambda x: len(x.split()))
    return " ".join(best.split()[:max_words])


# ---------------------------------------------------------------------------
# Web search + fetch
# ---------------------------------------------------------------------------

def search_web(query: str, max_results: int = 5, retries: int = 3) -> List[dict]:
    """Search the web. Retries with backoff (search engines rate-limit)."""
    for attempt in range(retries):
        try:
            with DDGS() as ddgs:
                return list(ddgs.text(query, max_results=max_results))
        except Exception as e:
            if "no results" in str(e).lower():
                return []
            wait = 2 * (2 ** attempt)
            print(f"  [warn] search error ({e}); retrying in {wait}s", file=sys.stderr)
            time.sleep(wait)
    return []


def fetch_page_text(url: str, timeout: int = 8, cache: Optional[dict] = None) -> Optional[str]:
    if cache is not None and url in cache:
        return cache[url]
    headers = {
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
        "Accept-Language": "en-US,en;q=0.9",
    }
    text = None
    try:
        resp = requests.get(url, headers=headers, timeout=timeout)
        resp.raise_for_status()
        if "text/html" in resp.headers.get("Content-Type", ""):
            text = clean_web_text(resp.text)
    except Exception:
        text = None  # caller falls back to the search snippet
    if cache is not None:
        cache[url] = text
    return text


def split_into_passages(text: str, window_sentences: int = 3) -> List[str]:
    """Break a fetched page into passages comparable in size to our chunks."""
    sentences = split_into_sentences(text)
    passages = []
    for i in range(0, len(sentences), window_sentences):
        passage = " ".join(sentences[i:i + window_sentences])
        if len(passage) > 20:
            passages.append(passage)
    return passages


# ---------------------------------------------------------------------------
# Core checking logic
# ---------------------------------------------------------------------------

def check_document(
    text: str,
    threshold: float = 0.80,
    results_per_chunk: int = 5,
    sentences_per_chunk: int = 3,
    model_name: str = "all-MiniLM-L6-v2",
    min_words: int = 8,
    delay: float = 1.0,
    verbose: bool = True,
) -> Report:
    _lazy_imports()

    if verbose:
        print(f"Loading embedding model '{model_name}' (first run downloads it)...")
    model = SentenceTransformer(model_name)

    all_chunks = chunk_document(text, sentences_per_chunk=sentences_per_chunk)
    chunks = [c for c in all_chunks if len(c.split()) >= min_words]
    report = Report(total_chunks=len(chunks), skipped_chunks=len(all_chunks) - len(chunks))
    page_cache: dict = {}

    if verbose:
        print(f"{len(chunks)} chunks to check ({report.skipped_chunks} too short, skipped).\n")

    for idx, chunk in enumerate(chunks):
        if verbose:
            print(f"[{idx + 1}/{len(chunks)}] {chunk[:70]}...")

        result = ChunkResult(chunk_index=idx, chunk_text=chunk)

        # Exact-phrase search first (best for copy-paste), then a looser one
        query = build_query(chunk)
        search_results = search_web(f'"{query}"', max_results=results_per_chunk)
        if not search_results:
            search_results = search_web(query, max_results=results_per_chunk)

        chunk_embedding = model.encode(chunk, convert_to_tensor=True)

        for sr in search_results:
            url = sr.get("href")
            if not url:
                continue

            page_text = fetch_page_text(url, cache=page_cache)
            if page_text:
                passages = split_into_passages(page_text, window_sentences=sentences_per_chunk)[:300]
            else:
                # Page blocked (403 etc.): fall back to the search-result snippet
                snippet = (sr.get("body") or "").strip()
                passages = [snippet] if len(snippet) > 20 else []
            if not passages:
                continue

            result.sources_checked += 1
            passage_embeddings = model.encode(passages, convert_to_tensor=True)
            cosine_scores = util.cos_sim(chunk_embedding, passage_embeddings)[0]

            best_idx = int(cosine_scores.argmax())
            best_score = float(cosine_scores[best_idx])

            if best_score > result.best_score:
                result.best_score = best_score
                result.best_source_url = url
                result.best_source_snippet = passages[best_idx][:200]

        if result.sources_checked == 0:
            report.no_result_chunks += 1
            if verbose:
                print("    -> no web sources could be compared")
        else:
            result.is_flagged = result.best_score >= threshold
            if result.is_flagged:
                report.flagged_chunks.append(result)
            if verbose:
                flag = "FLAGGED" if result.is_flagged else "ok"
                print(f"    -> best match {result.best_score:.3f} ({flag}): {result.best_source_url}")

        report.all_chunks.append(result)
        time.sleep(delay)  # be polite; avoids rate limiting

    return report


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------

def print_report(report: Report, threshold: float):
    print("\n" + "=" * 70)
    print("PLAGIARISM CHECK REPORT")
    print("=" * 70)
    print(f"Chunks checked       : {report.total_chunks} ({report.skipped_chunks} too short, skipped)")
    print(f"Web coverage         : {report.coverage:.0f}% of chunks had a source to compare")
    print(f"Flagged chunks       : {len(report.flagged_chunks)}")
    print(f"Flagged rate         : {report.overall_score:.1f}%")
    print(f"Similarity threshold : {threshold}")
    print("-" * 70)

    if report.coverage < 70:
        print("WARNING: low web coverage. Many chunks had no source to compare, so a low")
        print("flagged rate is NOT proof the text is original. Try a larger --delay")
        print("(e.g. 3) or run again later; the search engine may be rate-limiting.")
        print("-" * 70)

    if not report.flagged_chunks:
        print("No matches above threshold among the sources that were compared.")
    else:
        for r in sorted(report.flagged_chunks, key=lambda x: -x.best_score):
            print(f"\nChunk #{r.chunk_index + 1} (score: {r.best_score:.3f})")
            print(f"  Your text : {r.chunk_text[:150]}...")
            print(f"  Source    : {r.best_source_url}")
            print(f"  Matched   : {r.best_source_snippet}")

    print("\n" + "=" * 70)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Semantic web-based plagiarism checker")
    parser.add_argument("file", help="Path to the file to check (.txt, .docx or .pdf)")
    parser.add_argument("--threshold", type=float, default=0.80,
                         help="Similarity score (0-1) above which a chunk is flagged (default 0.80)")
    parser.add_argument("--results-per-chunk", type=int, default=5,
                         help="How many web search results to check per chunk (default 5)")
    parser.add_argument("--sentences-per-chunk", type=int, default=3,
                         help="How many sentences form one chunk (default 3)")
    parser.add_argument("--min-words", type=int, default=8,
                         help="Skip chunks shorter than this many words (default 8)")
    parser.add_argument("--delay", type=float, default=1.0,
                         help="Seconds to wait between chunks to avoid rate limits (default 1.0)")
    parser.add_argument("--keep-references", action="store_true",
                         help="Also check the References/Bibliography section (skipped by default)")
    parser.add_argument("--json", default=None,
                         help="Also write the results as JSON to this path "
                              "(consumed by plag_remover.py --report)")
    parser.add_argument("--model", default="all-MiniLM-L6-v2",
                         help="sentence-transformers model name (default all-MiniLM-L6-v2)")
    args = parser.parse_args()

    try:
        text = load_document(args.file)
    except OSError as e:
        print(f"Could not read file: {e}", file=sys.stderr)
        sys.exit(1)

    if not text.strip():
        print("No text could be extracted from the file.", file=sys.stderr)
        sys.exit(1)

    if not args.keep_references:
        text = strip_references(text)

    report = check_document(
        text,
        min_words=args.min_words,
        delay=args.delay,
        threshold=args.threshold,
        results_per_chunk=args.results_per_chunk,
        sentences_per_chunk=args.sentences_per_chunk,
        model_name=args.model,
    )
    print_report(report, args.threshold)

    if args.json:
        payload = {
            "total_chunks": report.total_chunks,
            "skipped_chunks": report.skipped_chunks,
            "no_result_chunks": report.no_result_chunks,
            "overall_score": report.overall_score,
            "coverage": report.coverage,
            "threshold": args.threshold,
            "flagged_chunks": [
                {
                    "chunk_index": r.chunk_index,
                    "chunk_text": r.chunk_text,
                    "best_score": r.best_score,
                    "best_source_url": r.best_source_url,
                    "best_source_snippet": r.best_source_snippet,
                }
                for r in report.flagged_chunks
            ],
        }
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2)
        print(f"JSON report written to {args.json}")


if __name__ == "__main__":
    main()