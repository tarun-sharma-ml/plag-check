import json

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from plagiarism_checker import load_document


def test_load_document_reads_utf8_text(tmp_path):
    path = tmp_path / "essay.txt"
    path.write_text("A document with café text.", encoding="utf-8")

    assert load_document(str(path)) == "A document with café text."


def test_load_document_falls_back_to_cp1252(tmp_path):
    path = tmp_path / "essay.txt"
    path.write_bytes(b"Text with caf\xe9.")

    assert load_document(str(path)) == "Text with café."


def test_load_document_reads_docx_paragraphs_and_tables(tmp_path):
    path = tmp_path / "essay.docx"
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("   ")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table text."
    table.cell(0, 1).text = " "
    document.save(path)

    assert load_document(str(path)) == "First paragraph.\nTable text."


def test_load_document_reads_pdf_text(tmp_path):
    path = tmp_path / "essay.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_ref}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Extracted PDF text.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as pdf_file:
        writer.write(pdf_file)

    assert load_document(str(path)).strip() == "Extracted PDF text."


def test_load_document_rejects_legacy_doc(tmp_path):
    path = tmp_path / "essay.doc"
    path.write_text("legacy document")

    with pytest.raises(SystemExit, match="Old .doc files are not supported"):
        load_document(str(path))


def test_load_document_reports_missing_text_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_document(str(tmp_path / "missing.txt"))


def test_report_json_has_expected_shape(monkeypatch, tmp_path):
    import sys

    import plagiarism_checker
    from plagiarism_checker import ChunkResult, Report

    report_path = tmp_path / "report.json"
    report = Report(
        total_chunks=1,
        flagged_chunks=[
            ChunkResult(
                chunk_index=0,
                chunk_text="Matched passage.",
                best_score=0.91,
                best_source_url="https://example.com",
                best_source_snippet="Source passage.",
                is_flagged=True,
                sources_checked=1,
            )
        ],
    )
    monkeypatch.setattr(
        sys,
        "argv",
        ["plagiarism_checker.py", "essay.txt", "--json", str(report_path)],
    )
    monkeypatch.setattr(plagiarism_checker, "load_document", lambda _: "Document text.")
    monkeypatch.setattr(plagiarism_checker, "check_document", lambda *_args, **_kwargs: report)
    monkeypatch.setattr(plagiarism_checker, "print_report", lambda *_args: None)

    plagiarism_checker.main()
    payload = json.loads(report_path.read_text(encoding="utf-8"))

    assert payload["total_chunks"] == 1
    assert payload["overall_score"] == 100.0
    assert payload["flagged_chunks"] == [
        {
            "chunk_index": 0,
            "chunk_text": "Matched passage.",
            "best_score": 0.91,
            "best_source_url": "https://example.com",
            "best_source_snippet": "Source passage.",
        }
    ]
